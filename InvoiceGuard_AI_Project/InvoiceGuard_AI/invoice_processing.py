import io
import os
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import pandas as pd

try:
    from sklearn.ensemble import IsolationForest
    ML_AVAILABLE = True
except ImportError:
    ML_AVAILABLE = False


CENT = Decimal("0.01")
_MONEY = (
    r"\(?-?(?:\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?|"
    r"\d[\d,]*(?:\.\d{1,2})?|\d+,\d{1,2})\)?"
)
_CURRENCY = r"(?:INR|USD|EUR|GBP|CAD|AUD|Rs\.?|₹|\$|€|£)?"
_NUMBER_RE = re.compile(_MONEY)
_SKIP_ITEM_LINE = re.compile(
    r"\b(?:sub\s*total|subtotal|grand\s*total|total\s*due|amount\s*due|"
    r"invoice\s*total|tax|gst|vat|discount|shipping|freight|balance|"
    r"invoice\s*(?:no|number|date)|due\s*date)\b",
    re.IGNORECASE,
)


def _money(value):
    if value is None:
        return None
    text = str(value).strip().replace("₹", "").replace("$", "").replace("€", "")
    text = text.replace("£", "").replace("Rs.", "").replace("Rs", "").strip()
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").strip()
    if "." in text and "," in text and re.search(r"\.\d{3},", text):
        text = text.replace(".", "").replace(",", ".")
    elif "," not in text and re.fullmatch(r"\d{1,3}(?:\.\d{3})+", text):
        text = text.replace(".", "")
    elif "," in text and "." not in text and re.fullmatch(r"\d+,\d{1,2}", text):
        text = text.replace(",", ".")
    else:
        text = text.replace(",", "")
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    return (-amount if negative else amount).quantize(CENT, rounding=ROUND_HALF_UP)


def _label_value(text, labels, value_pattern=r"([^\n]+)"):
    alternatives = "|".join(labels)
    match = re.search(
        rf"^\s*(?:{alternatives})\s*(?:[:#\-]\s*)?{value_pattern}",
        text,
        re.IGNORECASE | re.MULTILINE,
    )
    return match.group(1).strip(" \t:#") if match else None


def _label_money(text, labels):
    currency = rf"(?:\(\s*(?:INR|USD|EUR|GBP|CAD|AUD)\s*\)|{_CURRENCY})?"
    raw = _label_value(
        text,
        labels,
        rf"(?:{currency}\s*)?({_MONEY})",
    )
    return _money(raw)


def _confidence(label_value, fallback=False):
    if label_value:
        return 0.96
    return 0.62 if fallback else 0.0


def _extract_vendor(text):
    labeled = _label_value(
        text, [r"vendor", r"supplier", r"seller", r"issued\s+by"]
    )
    if labeled:
        return labeled, False

    ignored = re.compile(
        r"\b(?:tax\s+invoice|invoice|receipt|bill\s+to|ship\s+to|"
        r"purchase\s+order|statement)\b",
        re.IGNORECASE,
    )
    for line in text.splitlines()[:8]:
        candidate = line.strip(" \t|:-")
        if (
            len(candidate) >= 3
            and len(candidate) <= 100
            and ":" not in candidate
            and not ignored.search(candidate)
            and not re.search(r"\d{3,}", candidate)
        ):
            return candidate, True
    return None, False


def _extract_date(text, labels):
    value = _label_value(
        text,
        labels,
        r"(\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}|"
        r"[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4}|"
        r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})",
    )
    return value, bool(value)


def _extract_line_items(text):
    items = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or _SKIP_ITEM_LINE.search(stripped):
            continue

        # Handle an explicitly listed set of amounts without inventing quantities.
        plus_values = re.findall(rf"(?:₹|\$|€|£)?\s*({_MONEY})", stripped)
        if "+" in stripped and len(plus_values) >= 2:
            prefix = re.split(rf"(?:₹|\$|€|£)?\s*{_MONEY}", stripped, maxsplit=1)[0]
            description = re.sub(r"[\s:|+\-]+$", "", prefix).strip()
            for index, value in enumerate(plus_values, start=1):
                items.append(
                    {
                        "description": description or f"Line item {index}",
                        "quantity": None,
                        "unit_price": None,
                        "line_total": _money(value),
                    }
                )
            continue

        # Rows should provide quantity, unit price, and extended amount.
        numbers = list(_NUMBER_RE.finditer(stripped))
        if len(numbers) < 3:
            continue
        quantity_match, unit_match, total_match = numbers[-3:]
        quantity = _money(quantity_match.group())
        unit_price = _money(unit_match.group())
        line_total = _money(total_match.group())
        if quantity is None or unit_price is None or line_total is None or quantity <= 0:
            continue
        description = stripped[: quantity_match.start()].strip(" \t|:-#.")
        if not description or len(description) > 120:
            continue
        items.append(
            {
                "description": description,
                "quantity": quantity,
                "unit_price": unit_price,
                "line_total": line_total,
            }
        )
    return items


def _words_number(words):
    small = {
        "zero": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5,
        "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
        "eleven": 11, "twelve": 12, "thirteen": 13, "fourteen": 14,
        "fifteen": 15, "sixteen": 16, "seventeen": 17, "eighteen": 18,
        "nineteen": 19,
    }
    tens = {
        "twenty": 20, "thirty": 30, "forty": 40, "fifty": 50,
        "sixty": 60, "seventy": 70, "eighty": 80, "ninety": 90,
    }
    scales = {
        "hundred": 100, "thousand": 1000, "lakh": 100000,
        "lac": 100000, "million": 1000000, "crore": 10000000,
    }
    normalized = re.sub(r"[^a-z ]", " ", words.lower())
    ignored = {"only", "rupee", "rupees", "inr", "dollar", "dollars", "euro", "euros", "pound", "pounds"}
    tokens = [word for word in normalized.split() if word not in ignored]

    def parse_integer(parts):
        total = 0
        group = 0
        found = False
        for word in parts:
            if word in small:
                group += small[word]
                found = True
            elif word in tens:
                group += tens[word]
                found = True
            elif word in scales:
                scale = scales[word]
                group = max(1, group)
                if scale == 100:
                    group *= scale
                else:
                    total += group * scale
                    group = 0
                found = True
            elif word == "and":
                continue
            else:
                return None
        return total + group if found else 0

    subunit_index = next(
        (index for index, word in enumerate(tokens) if word in {"paise", "cents"}),
        None,
    )
    if subunit_index is not None:
        before = tokens[:subunit_index]
        if "and" in before:
            split = len(before) - 1 - before[::-1].index("and")
            whole = parse_integer(before[:split])
            subunits = parse_integer(before[split + 1:])
        else:
            whole = parse_integer(before)
            subunits = 0
        if whole is None or subunits is None or subunits > 99:
            return None
        return Decimal(whole) + Decimal(subunits) / Decimal("100")
    value = parse_integer(tokens)
    return Decimal(value) if value else (Decimal("0") if "zero" in tokens else None)


def validate_calculations(fields):
    results = []
    items = fields["line_items"]
    if items and any(
        item["quantity"] is None or item["unit_price"] is None
        for item in items
    ):
        results.append(
            {
                "status": "uncertain",
                "message": "Quantity or unit price was not extracted for every line item; those extensions cannot be independently verified.",
            }
        )
    computable = [
        item for item in items
        if item["quantity"] is not None and item["unit_price"] is not None
        and item["line_total"] is not None
    ]
    for item in computable:
        expected = (item["quantity"] * item["unit_price"]).quantize(
            CENT, rounding=ROUND_HALF_UP
        )
        if abs(expected - item["line_total"]) > CENT:
            results.append(
                {
                    "status": "mismatch",
                    "message": (
                        f"Confirmed arithmetic mismatch for {item['description']}: "
                        f"quantity x unit price = {expected:.2f}, but the line total "
                        f"is {item['line_total']:.2f}."
                    ),
                }
            )
    if items and fields["subtotal"] is not None and fields["adjustment_extraction_uncertain"]:
        results.append(
            {
                "status": "uncertain",
                "message": "Line-item subtotal check skipped because a stated discount or charge could not be extracted.",
            }
        )
    elif items and fields["subtotal"] is not None:
        if computable and len(computable) == len(items):
            item_sum = sum(
                (item["quantity"] * item["unit_price"] for item in computable),
                Decimal("0"),
            )
            check_name = "Calculated line items"
        elif all(item["line_total"] is not None for item in items):
            item_sum = sum(
                (item["line_total"] for item in items), Decimal("0")
            )
            check_name = "Stated line items"
        else:
            item_sum = None
            check_name = "Line items"
        if item_sum is not None:
            item_sum = (
                item_sum - (fields["discount"] or Decimal("0"))
                + (fields["charges"] or Decimal("0"))
            ).quantize(CENT, rounding=ROUND_HALF_UP)
            difference = abs(item_sum - fields["subtotal"])
            results.append(
                {
                    "status": "pass" if difference <= CENT else "mismatch",
                    "message": (
                        f"{check_name} sum to {item_sum:.2f}; stated subtotal is "
                        f"{fields['subtotal']:.2f}."
                    ),
                }
            )
    elif items:
        results.append(
            {"status": "uncertain", "message": "Line-item subtotal check skipped because subtotal is missing."}
        )

    subtotal, tax, total = fields["subtotal"], fields["tax"], fields["total"]
    if fields["adjustment_extraction_uncertain"]:
        results.append(
            {
                "status": "uncertain",
                "message": "Grand-total check skipped because a stated discount or charge could not be extracted.",
            }
        )
    elif subtotal is not None and total is not None and (tax is not None or fields["tax_rate"] is not None):
        tax_amount = tax
        if tax_amount is None and fields["tax_rate"] is not None:
            tax_amount = (subtotal * fields["tax_rate"] / Decimal("100")).quantize(
                CENT, rounding=ROUND_HALF_UP
            )
        discount = fields["discount"] if fields["discount"] is not None else Decimal("0")
        charges = fields["charges"] if fields["charges"] is not None else Decimal("0")
        expected = (subtotal - discount + tax_amount + charges).quantize(
            CENT, rounding=ROUND_HALF_UP
        )
        adjustment_note = (
            " No separate discounts or charges were detected."
            if fields["discount"] is None and fields["charges"] is None
            else ""
        )
        results.append(
            {
                "status": "pass" if abs(expected - total) <= CENT else "mismatch",
                "message": (
                    f"Subtotal - stated discount + stated tax + stated charges = {expected:.2f}; "
                    f"stated grand total is {total:.2f}.{adjustment_note}"
                ),
            }
        )
    else:
        results.append(
            {
                "status": "uncertain",
                "message": "Grand-total calculation could not be confirmed because required amounts or tax details are missing.",
            }
        )

    rate = fields["tax_rate"]
    if subtotal is not None and tax is not None and rate is not None:
        expected_tax = (subtotal * rate / Decimal("100")).quantize(
            CENT, rounding=ROUND_HALF_UP
        )
        results.append(
            {
                "status": "pass" if abs(expected_tax - tax) <= CENT else "mismatch",
                "message": (
                    f"Tax at {rate}% on subtotal is {expected_tax:.2f}; "
                    f"stated tax is {tax:.2f}."
                ),
            }
        )

    words_value = fields["amount_in_words"]
    if words_value and total is not None:
        parsed_words = _words_number(words_value)
        if parsed_words is None:
            results.append(
                {
                    "status": "uncertain",
                    "message": "Amount in words was found but could not be reliably interpreted.",
                }
            )
        else:
            results.append(
                {
                    "status": "pass" if Decimal(parsed_words).quantize(CENT) == total else "mismatch",
                    "message": (
                        f"Amount in words represents {parsed_words:.2f}; "
                        f"numeric grand total is {total:.2f}."
                    ),
                }
            )
    return results


def extract_fields(text, filename):
    vendor, vendor_fallback = _extract_vendor(text)
    invoice_number = _label_value(
        text,
        [r"invoice\s*(?:no\.?|number|#|id)", r"bill\s*(?:no\.?|number|#)"],
        r"([A-Z0-9][A-Z0-9\-\/]{1,})",
    )
    invoice_date, _ = _extract_date(
        text, [r"invoice\s*date", r"issued\s*on", r"date"]
    )
    due_date, _ = _extract_date(text, [r"due\s*date", r"payment\s*due"])
    customer = _label_value(
        text, [r"bill\s*to", r"billed\s*to", r"customer", r"buyer", r"client"]
    )
    subtotal = _label_money(text, [r"sub\s*total", r"subtotal", r"net\s*amount"])
    tax = _label_money(
        text,
        [
            r"(?:gst|vat|sales\s*tax|tax)(?:\s*amount)?"
            r"(?!\s*(?:rate|percentage))"
            r"(?:\s*\(\s*\d+(?:\.\d+)?%\s*\)|\s*@\s*\d+(?:\.\d+)?%|"
            r"\s+at\s+\d+(?:\.\d+)?%|\s+\d+(?:\.\d+)?%)?"
        ],
    )
    if tax is None:
        tax = _label_money(
            text, [r"(?:cgst|sgst|igst)(?:\s*amount)?(?:\s*\(\s*\d+(?:\.\d+)?%\s*\))?"]
        )
    total = _label_money(
        text,
        [r"grand\s*total", r"invoice\s*total", r"total\s*amount", r"total\s*due", r"total"],
    )
    if total is None:
        total = _label_money(text, [r"amount\s*due", r"amount\s*payable"])
    amount_due = _label_money(text, [r"amount\s*due", r"balance\s*due"])
    tax_rate_match = re.search(
        r"\b(?:GST|VAT|tax|CGST|SGST|IGST)\s+(?:rate|percentage)\s*[:\-]?\s*(\d+(?:\.\d+)?)\s*%",
        text,
        re.IGNORECASE,
    ) or re.search(
        r"\b(?:GST|VAT|tax|CGST|SGST|IGST)\b\s*(?:\(\s*|@\s*|at\s*|\s*)(\d+(?:\.\d+)?)\s*%",
        text,
        re.IGNORECASE,
    ) or re.search(
        r"(\d+(?:\.\d+)?)\s*%\s*(?:GST|VAT|tax|CGST|SGST|IGST)\b",
        text,
        re.IGNORECASE,
    )
    tax_rate = Decimal(tax_rate_match.group(1)) if tax_rate_match else None
    discount = _label_money(text, [r"discount", r"less"])
    charges = _label_money(text, [r"shipping(?:\s+charges?)?", r"freight", r"other\s+charges?"])
    adjustment_extraction_uncertain = (
        discount is None
        and bool(re.search(r"^\s*(?:discount|less)\b", text, re.IGNORECASE | re.MULTILINE))
    ) or (
        charges is None
        and bool(re.search(r"^\s*(?:shipping|freight|other\s+charges?)\b", text, re.IGNORECASE | re.MULTILINE))
    )
    amount_in_words = _label_value(
        text, [r"amount\s+in\s+words", r"total\s+in\s+words", r"rupees\s+in\s+words"]
    )
    currency_match = re.search(r"\b(INR|USD|EUR|GBP|CAD|AUD)\b", text, re.IGNORECASE)
    if currency_match:
        currency = currency_match.group(1).upper()
    elif "₹" in text or re.search(r"\bRs\.?\s*\d", text, re.IGNORECASE):
        currency = "INR"
    elif "$" in text:
        currency = "USD"
    elif "€" in text:
        currency = "EUR"
    elif "£" in text:
        currency = "GBP"
    else:
        currency = None

    field_confidence = {
        "vendor": _confidence(vendor, vendor_fallback),
        "invoice_number": _confidence(invoice_number),
        "invoice_date": _confidence(invoice_date),
        "due_date": _confidence(due_date),
        "customer": _confidence(customer),
        "subtotal": _confidence(subtotal),
        "tax": _confidence(tax),
        "total": _confidence(total),
    }
    measured = [value for value in field_confidence.values() if value > 0]
    fields = {
        "vendor": vendor,
        "invoice_number": invoice_number,
        "invoice_date": invoice_date,
        "due_date": due_date,
        "customer": customer,
        "currency": currency,
        "line_items": _extract_line_items(text),
        "subtotal": subtotal,
        "tax_rate": tax_rate,
        "tax": tax,
        "discount": discount,
        "charges": charges,
        "adjustment_extraction_uncertain": adjustment_extraction_uncertain,
        "total": total,
        "amount_due": amount_due,
        "amount_in_words": amount_in_words,
        "field_confidence": field_confidence,
        "confidence": round(sum(measured) / len(measured), 2) if measured else 0.0,
        "file_name": filename,
        "extracted_text": text,
    }
    fields["validation_results"] = validate_calculations(fields)
    return fields


def evaluate_invoice(fields, file_hash, existing_records):
    findings = []
    adjustments = []
    score = 0

    def adjust(reason, points):
        nonlocal score
        applied_points = min(points, 100 - score)
        if applied_points > 0:
            score += applied_points
            adjustments.append({"reason": reason, "points": applied_points})

    if not fields["vendor"]:
        findings.append("Extraction uncertainty: vendor could not be confidently extracted.")
        adjust("Vendor is missing; manual review is needed.", 10)
    if not fields["invoice_number"]:
        findings.append("Extraction uncertainty: invoice number is missing or could not be extracted.")
        adjust("Invoice number is missing; duplicate checks are less reliable.", 5)
    if fields["total"] is None:
        findings.append("Extraction uncertainty: grand total could not be confidently extracted.")
        adjust("Grand total is missing; calculations cannot be confirmed.", 20)

    for name, confidence in fields["field_confidence"].items():
        if confidence and confidence < 0.75:
            findings.append(f"Extraction uncertainty: {name.replace('_', ' ')} was inferred with lower confidence.")
    mismatches = [
        result for result in fields["validation_results"]
        if result["status"] == "mismatch"
    ]
    for result in mismatches:
        findings.append(f"Confirmed arithmetic mismatch: {result['message']}")
        adjust(result["message"], 20)
    uncertainties = [
        result for result in fields["validation_results"]
        if result["status"] == "uncertain"
    ]
    for result in uncertainties:
        findings.append(f"Calculation uncertainty: {result['message']}")

    if fields["total"] is not None and fields["total"] < 0:
        findings.append("Negative total amount detected.")
        adjust("Invoice total is negative and requires review.", 30)
    if fields["total"] is not None and fields["total"] == 0:
        findings.append("Total amount is zero; verify the invoice.")
        adjust("Invoice total is zero and requires review.", 15)

    for row in existing_records:
        if row.get("file_hash") == file_hash:
            findings.append("Exact file duplicate: this document has already been uploaded.")
            adjust("Exact file duplicate detected.", 70)
            break
    for row in existing_records:
        vendor = (fields["vendor"] or "").strip().lower()
        invoice_number = (fields["invoice_number"] or "").strip().lower()
        same_vendor = str(row.get("vendor") or "").strip().lower() == vendor and bool(vendor)
        same_number = (
            str(row.get("invoice_number") or "").strip().lower() == invoice_number
            and bool(invoice_number)
        )
        row_total = row.get("total")
        same_total = (
            row_total is not None
            and not pd.isna(row_total)
            and fields["total"] is not None
            and abs(float(row_total) - float(fields["total"])) < 0.01
        )
        same_date = (row.get("invoice_date") or "") == (fields["invoice_date"] or "")
        if same_vendor and (same_number or (same_total and same_date)):
            if not any("Possible duplicate:" in finding for finding in findings):
                findings.append("Possible duplicate: vendor and invoice details closely match a previously uploaded record.")
                adjust("Possible duplicate invoice details detected.", 55)
            break

    if (
        ML_AVAILABLE
        and fields["total"] is not None
        and fields["subtotal"] is not None
        and fields["tax"] is not None
    ):
        numeric_records = [
            {
                key: row.get(key)
                for key in ("subtotal", "tax", "total")
            }
            for row in existing_records
        ]
        numeric = pd.DataFrame(numeric_records).apply(pd.to_numeric, errors="coerce").dropna()
        if len(numeric) >= 10:
            current = pd.DataFrame([{
                "subtotal": float(fields["subtotal"]),
                "tax": float(fields["tax"]),
                "total": float(fields["total"]),
            }])
            model = IsolationForest(contamination=0.1, random_state=42)
            model.fit(numeric)
            if model.predict(current)[0] == -1:
                findings.append("ML anomaly signal: amount pattern differs from historical invoices. This is not proof of fraud.")
                adjust("Amount pattern differs from historical invoices; this is not proof of fraud.", 15)

    critical_missing = not fields["vendor"] or fields["total"] is None
    if critical_missing:
        status = "Needs review - critical fields missing"
    elif mismatches:
        status = "Review recommended - calculation mismatch"
    elif score >= 60:
        status = "High review priority"
    elif score >= 15 or uncertainties or any(
        finding.startswith("Extraction uncertainty:") for finding in findings
    ):
        status = "Review recommended"
    elif findings:
        status = "Minor warning"
    else:
        status = "No issue detected"
    if not findings:
        findings = ["No configured checks were triggered. This does not guarantee the invoice is legitimate."]
    return score, status, findings, adjustments


def _ocr_image(image):
    try:
        import pytesseract
        from PIL import ImageEnhance, ImageFilter, ImageOps
    except ImportError as exc:
        raise RuntimeError("Image OCR requires the pytesseract and Pillow packages.") from exc
    tesseract_cmd = os.environ.get("TESSERACT_CMD")
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    image = ImageOps.exif_transpose(image).convert("L")
    image = ImageEnhance.Contrast(image).enhance(1.7).filter(ImageFilter.SHARPEN)
    try:
        return pytesseract.image_to_string(image, config="--psm 6")
    except Exception as exc:
        raise RuntimeError(
            "Tesseract OCR failed. Install Tesseract and add it to PATH or set TESSERACT_CMD."
        ) from exc


def read_upload(upload):
    data = upload.getvalue()
    name = upload.name
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        try:
            import pymupdf
            with pymupdf.open(stream=data, filetype="pdf") as document:
                text = "\n".join(page.get_text("text") for page in document)
        except Exception as exc:
            return "", data, f"pdf-error: Could not read PDF ({exc})."
        if len(re.sub(r"\W", "", text, flags=re.UNICODE)) >= 8:
            return text, data, "pdf-text"
        try:
            from pdf2image import convert_from_bytes
        except ImportError:
            return "", data, (
                "pdf-ocr-unavailable: PDF has no usable text; install pdf2image "
                "and Poppler to OCR scanned PDFs."
            )
        try:
            pages = convert_from_bytes(
                data, dpi=220, poppler_path=os.environ.get("POPPLER_PATH") or None
            )
            text = "\n".join(_ocr_image(page) for page in pages)
        except Exception as exc:
            return "", data, (
                "pdf-ocr-error: Scanned PDF OCR failed. Ensure Poppler is on PATH "
                f"or set POPPLER_PATH; check Tesseract as well ({exc})."
            )
        if not text.strip():
            return "", data, "pdf-ocr-empty: OCR completed but found no readable text."
        return text, data, "pdf-ocr"
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        try:
            from PIL import Image
            with Image.open(io.BytesIO(data)) as image:
                text = _ocr_image(image)
        except Exception as exc:
            return "", data, f"image-error: Could not OCR image ({exc})."
        return text, data, "image-ocr"
    if suffix in {".txt", ".csv"}:
        return data.decode("utf-8-sig", errors="replace"), data, "text"
    if suffix == ".docx":
        try:
            from docx import Document
            document = Document(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            paragraphs.extend(
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            )
            text = "\n".join(paragraphs)
        except ImportError:
            return "", data, "docx-error: Install python-docx to read DOCX invoices."
        except Exception as exc:
            return "", data, f"docx-error: Could not read DOCX ({exc})."
        if not text.strip():
            return "", data, "docx-empty: No readable text found in this document."
        return text, data, "docx-text"
    return "", data, f"unsupported: Unsupported file type '{suffix or 'unknown'}'."
