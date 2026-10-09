import io
import unittest
from decimal import Decimal

import pymupdf
from docx import Document

from invoice_processing import evaluate_invoice, extract_fields, read_upload


INCORRECT_INVOICE = """Northstar Office Supplies (Fictional)
Vendor: Northstar Office Supplies (Fictional)
Invoice Number: NS-2026-001
Invoice Date: 2026-10-01
Due Date: 2026-10-31
Bill To: Example Customer
Description | Qty | Unit Price | Line Total
Paper | 1 | INR 1350 | INR 1350
Pens | 1 | INR 400 | INR 400
Folders | 2 | INR 250 | INR 600
Subtotal: INR 2350
GST (18%): INR 423
Grand Total: INR 2700
"""

CORRECT_INVOICE = """Northstar Office Supplies (Fictional)
Vendor: Northstar Office Supplies (Fictional)
Invoice Number: NS-2026-002
Invoice Date: 2026-10-01
Description | Qty | Unit Price | Line Total
Paper | 1 | INR 1350 | INR 1350
Pens | 1 | INR 400 | INR 400
Folders | 2 | INR 250 | INR 500
Subtotal: INR 2250
GST (18%): INR 405
Grand Total: INR 2655
"""


class MemoryUpload:
    def __init__(self, name, data):
        self.name = name
        self._data = data

    def getvalue(self):
        return self._data


class InvoiceExtractionTests(unittest.TestCase):
    def test_incorrect_invoice_extracts_fields_and_flags_three_calculations(self):
        fields = extract_fields(INCORRECT_INVOICE, "incorrect.txt")
        self.assertEqual(fields["vendor"], "Northstar Office Supplies (Fictional)")
        self.assertEqual(fields["total"], Decimal("2700.00"))
        self.assertEqual(fields["invoice_number"], "NS-2026-001")
        self.assertEqual(fields["due_date"], "2026-10-31")
        self.assertEqual(fields["customer"], "Example Customer")
        self.assertEqual(fields["tax_rate"], Decimal("18"))

        mismatches = [
            result["message"]
            for result in fields["validation_results"]
            if result["status"] == "mismatch"
        ]
        self.assertEqual(len(mismatches), 3)
        self.assertTrue(any("Folders" in message for message in mismatches))
        self.assertTrue(any("subtotal is 2350.00" in message for message in mismatches))
        self.assertTrue(any("grand total is 2700.00" in message for message in mismatches))

        score, status, findings, adjustments = evaluate_invoice(fields, "hash-a", [])
        self.assertGreater(score, 0)
        self.assertIn("calculation mismatch", status)
        self.assertTrue(any("Confirmed arithmetic mismatch" in item for item in findings))
        self.assertEqual(len(adjustments), 3)

    def test_correct_invoice_passes_all_calculations_and_scores_zero(self):
        fields = extract_fields(CORRECT_INVOICE, "correct.txt")
        self.assertEqual(fields["vendor"], "Northstar Office Supplies (Fictional)")
        self.assertEqual(fields["total"], Decimal("2655.00"))
        self.assertEqual(fields["line_items"][2]["quantity"], Decimal("2.00"))
        self.assertEqual(fields["line_items"][2]["unit_price"], Decimal("250.00"))
        self.assertEqual(fields["line_items"][2]["line_total"], Decimal("500.00"))
        self.assertTrue(
            all(result["status"] == "pass" for result in fields["validation_results"])
        )
        score, status, _, adjustments = evaluate_invoice(fields, "hash-b", [])
        self.assertEqual(score, 0)
        self.assertEqual(status, "No issue detected")
        self.assertEqual(adjustments, [])

    def test_duplicate_detection_adds_an_explained_score_adjustment(self):
        fields = extract_fields(CORRECT_INVOICE, "correct.txt")
        existing = [{
            "file_hash": "previous-hash",
            "vendor": fields["vendor"],
            "invoice_number": fields["invoice_number"],
            "invoice_date": fields["invoice_date"],
            "total": float(fields["total"]),
            "subtotal": float(fields["subtotal"]),
            "tax": float(fields["tax"]),
        }]
        score, status, findings, adjustments = evaluate_invoice(
            fields, "new-hash", existing
        )
        self.assertEqual(score, 55)
        self.assertEqual(status, "Review recommended")
        self.assertTrue(any("Possible duplicate" in finding for finding in findings))
        self.assertEqual(adjustments, [{
            "reason": "Possible duplicate invoice details detected.",
            "points": 55,
        }])

    def test_total_label_does_not_match_subtotal_and_missing_values_stay_none(self):
        fields = extract_fields(
            "Vendor: Example Supplier\nSubtotal: 100.00\n", "partial.txt"
        )
        self.assertEqual(fields["vendor"], "Example Supplier")
        self.assertEqual(fields["subtotal"], Decimal("100.00"))
        self.assertIsNone(fields["total"])
        score, status, findings, adjustments = evaluate_invoice(fields, "hash-c", [])
        self.assertGreater(score, 0)
        self.assertIn("critical fields missing", status)
        self.assertTrue(any("grand total" in item for item in findings))
        self.assertTrue(any(item["points"] > 0 for item in adjustments))

    def test_amount_due_and_words_are_checked_separately(self):
        fields = extract_fields(
            "Vendor: Example Supplier\n"
            "Subtotal: INR 150.00\n"
            "Tax Amount: INR 0.00\n"
            "Grand Total: INR 150.00\n"
            "Amount Due: INR 100.00\n"
            "Amount in words: Rupees One Hundred Fifty Only\n",
            "due.txt",
        )
        self.assertEqual(fields["total"], Decimal("150.00"))
        self.assertEqual(fields["amount_due"], Decimal("100.00"))
        words_check = [
            result for result in fields["validation_results"]
            if "Amount in words" in result["message"]
        ]
        self.assertEqual(words_check[0]["status"], "pass")
        paise_fields = extract_fields(
            "Vendor: Example Supplier\n"
            "Subtotal: INR 150.50\n"
            "Tax Amount: INR 0.00\n"
            "Grand Total: INR 150.50\n"
            "Amount in words: Rupees One Hundred Fifty and Fifty Paise Only\n",
            "paise.txt",
        )
        paise_check = [
            result for result in paise_fields["validation_results"]
            if "Amount in words" in result["message"]
        ]
        self.assertEqual(paise_check[0]["status"], "pass")

    def test_currency_symbols_grouped_amounts_and_tax_rate_label(self):
        fields = extract_fields(
            "Northstar Supplies\n"
            "Invoice No: NS/5\n"
            "Invoice Date: 01/10/2026\n"
            "Subtotal: ₹2,350.00\n"
            "GST Rate: 18%\n"
            "Tax Amount: ₹423.00\n"
            "Grand Total: ₹2,773.00\n",
            "symbols.txt",
        )
        self.assertEqual(fields["vendor"], "Northstar Supplies")
        self.assertEqual(fields["invoice_number"], "NS/5")
        self.assertEqual(fields["currency"], "INR")
        self.assertEqual(fields["subtotal"], Decimal("2350.00"))
        self.assertEqual(fields["tax_rate"], Decimal("18"))
        self.assertEqual(fields["tax"], Decimal("423.00"))
        self.assertEqual(fields["total"], Decimal("2773.00"))

    def test_decimal_comma_and_period_grouping(self):
        fields = extract_fields(
            "Vendor: Example Supplier\n"
            "Subtotal: EUR 1.234,56\n"
            "Tax Amount: EUR 0,00\n"
            "Grand Total: EUR 1.234,56\n",
            "european.txt",
        )
        self.assertEqual(fields["subtotal"], Decimal("1234.56"))
        self.assertEqual(fields["tax"], Decimal("0.00"))
        self.assertEqual(fields["total"], Decimal("1234.56"))
        self.assertEqual(fields["currency"], "EUR")

    def test_txt_and_text_pdf_upload_readers(self):
        text_upload = MemoryUpload("invoice.txt", b"Vendor: Example\nGrand Total: 12")
        text, _, method = read_upload(text_upload)
        self.assertEqual(method, "text")
        self.assertIn("Grand Total", text)

        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Vendor: PDF Supplier\nGrand Total: 125.00")
        pdf_data = document.tobytes()
        document.close()
        pdf_upload = MemoryUpload("invoice.pdf", pdf_data)
        text, _, method = read_upload(pdf_upload)
        self.assertEqual(method, "pdf-text")
        self.assertIn("PDF Supplier", text)

    def test_docx_upload_reads_paragraph_and_table_text(self):
        document = Document()
        document.add_paragraph("Vendor: DOCX Supplier")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Grand Total"
        table.cell(0, 1).text = "INR 250.00"
        buffer = io.BytesIO()
        document.save(buffer)
        upload = MemoryUpload("invoice.docx", buffer.getvalue())

        text, _, method = read_upload(upload)
        self.assertEqual(method, "docx-text")
        self.assertIn("DOCX Supplier", text)
        self.assertIn("INR 250.00", text)

    def test_unknown_file_type_has_actionable_error(self):
        text, _, method = read_upload(MemoryUpload("invoice.xlsx", b""))
        self.assertEqual(text, "")
        self.assertIn("Unsupported file type", method)

    def test_corrupt_pdf_and_image_return_reader_errors(self):
        pdf_text, _, pdf_method = read_upload(MemoryUpload("broken.pdf", b"not a PDF"))
        image_text, _, image_method = read_upload(MemoryUpload("broken.png", b"not an image"))
        self.assertEqual(pdf_text, "")
        self.assertTrue(pdf_method.startswith("pdf-error:"))
        self.assertEqual(image_text, "")
        self.assertTrue(image_method.startswith("image-error:"))


if __name__ == "__main__":
    unittest.main()
